"""API integration tests for Document Dossier Word Export endpoint (Bloque 12 — Fase 4B).

Covers all required verification scenarios:
1. POST correcto con 1 documento -> HTTP 200.
2. POST correcto con varios documentos -> HTTP 200.
3. Se respeta el orden recibido por el cliente (autoridad del orden de entrada).
4. Título personalizado en portada y filename.
5. Sin título usa título por defecto institucional.
6. Documento sin análisis se procesa y exporta correctamente.
7. Dossier mixto (analizados y no analizados) exporta correctamente.
8. Documento inexistente -> HTTP 404 con identificación del ID ausente.
9. IDs duplicados en request -> rechazo con HTTP 422.
10. Lista vacía de IDs -> HTTP 422 ("al menos 1 documento").
11. Más de 30 IDs -> HTTP 422 ("límite máximo de 30 documentos").
12. UUID inválido -> HTTP 422.
13. Usuario no autenticado -> HTTP 401.
14. Content-Type correcto (application/vnd.openxmlformats-officedocument.wordprocessingml.document).
15. Content-Disposition presente y seguro.
16. Contenido descargado es un ZIP/DOCX binario válido.
17. El DOCX contiene realmente los documentos solicitados.
18. El orden dentro del DOCX coincide con el request.
19. El título aparece correctamente en los párrafos de portada.
20. No se generan archivos temporales en disco.
21. No se llama a Gemini ni a proveedores de IA.
22. No se modifica la base de datos (garantía read-only).
23. El endpoint individual de exportación sigue funcionando sin regresión.
24. Tests de seguridad: prevención de path traversal, inyección CRLF en Content-Disposition y acceso no autenticado.
25. Verificación de registro en OpenAPI (/openapi.json).
"""

from __future__ import annotations

import io
import os
import tempfile
import unittest.mock as mock
import uuid
from datetime import datetime, timezone, timedelta
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import select, func
from sqlalchemy.orm import Session
import docx

from app.main import app
from app.api.v1.endpoints.auth import get_current_user
from app.models.entry import Entry
from app.models.source import Source, SourceType
from app.models.tracking import TrackedEntity, TrackingMatrix, TrackingTopic
from app.models.analysis import AnalysisPromptVersion, EntryAnalysis, EntryAnalysisTopic, AnalysisCall
from app.api.v1.endpoints.documents import DOCX_MIME_TYPE
from app.services.analysis_service import compute_analysis_input_hash


@pytest.fixture
def setup_dossier_data(db_session: Session):
    """Seed sample analysed and unanalysed documentary records for dossier export tests."""
    matrix = TrackingMatrix(
        code="MATRIX-DOSSIER-12",
        name="Matriz Dossier Export",
        status="active",
    )
    db_session.add(matrix)
    db_session.flush()

    topic1 = TrackingTopic(
        matrix_id=matrix.id,
        code="antitrust_cartels",
        name="Cárteles y Colusión",
    )
    topic2 = TrackingTopic(
        matrix_id=matrix.id,
        code="vertical_agreements",
        name="Acuerdos Verticales",
    )
    db_session.add_all([topic1, topic2])
    db_session.flush()

    entity = TrackedEntity(
        display_name="Comisión Nacional de los Mercados y la Competencia",
        entity_type="institution",
    )
    db_session.add(entity)
    db_session.flush()

    source = Source(
        name="Sede Electrónica CNMC",
        url="https://cnmc.example.org",
        type=SourceType.WEBSITE,
        category="official",
        tracked_entity_id=entity.id,
    )
    db_session.add(source)
    db_session.flush()

    now = datetime.now(timezone.utc)

    prompt_version = AnalysisPromptVersion(
        code="prompt_dossier_test_v1",
        version=1,
        stage="deep_analysis",
        name="Deep Analysis Dossier Prompt",
        system_prompt="System instructions",
        user_prompt_template="User template",
        response_schema_version="v1",
    )
    db_session.add(prompt_version)
    db_session.flush()

    # Document 1: Fully analysed cartel decision
    entry_1 = Entry(
        source_id=source.id,
        title="Resolución CNMC S/0010/24 sobre Cártel de Fabricantes de Transformadores",
        url="https://cnmc.example.org/res-transf-0010",
        canonical_url="https://cnmc.example.org/res-transf-0010",
        content=(
            "La Dirección de Competencia de la CNMC sanciona a cinco mercantiles por acordar precios "
            "y repartirse las licitaciones convocadas por empresas de distribución eléctrica. "
            "Queda acreditada la concurrencia continuada y deliberada para falsear la competencia en el mercado nacional."
        ),
        excerpt="Sanción por cártel en licitaciones de transformadores...",
        author="Consejo de la CNMC",
        published_at=now - timedelta(days=10),
        captured_at=now - timedelta(days=10),
        content_type="judicial_decision",
        language="es",
        raw_metadata={"expediente": "S/0010/24"},
    )
    db_session.add(entry_1)
    db_session.flush()

    content_hash_1 = compute_analysis_input_hash(entry_1)
    analysis_1 = EntryAnalysis(
        entry_id=entry_1.id,
        matrix_id=matrix.id,
        pipeline_version="v7",
        status="completed",
        entry_content_hash=content_hash_1,
        matrix_snapshot={"topics": [{"code": "antitrust_cartels", "name": "Cárteles y Colusión"}]},
        matrix_snapshot_hash="hash-dossier-1",
        relevance_status="relevant",
        relevance_score=96,
        confidence=0.95,
        summary="Acreditación de cártel colusorio en licitaciones de transformadores con imposición de multas millonarias.",
        key_points=[
            "Infracción por objeto tipificada en el artículo 1 de la Ley 15/2007.",
            "Desestimación de las excepciones de prescripción presentadas por las empresas sancionadas.",
        ],
        completed_at=now - timedelta(days=9),
    )
    db_session.add(analysis_1)
    db_session.flush()

    eat_1 = EntryAnalysisTopic(analysis_id=analysis_1.id, topic_id=topic1.id, is_primary=True)
    call_1 = AnalysisCall(
        entry_analysis_id=analysis_1.id,
        prompt_version_id=prompt_version.id,
        stage="deep_analysis",
        provider="gemini_api",
        model="gemini-3.8-flash",
        status="completed",
        raw_response={
            "result": {
                "summary_evidence": [
                    {
                        "source_field": "content",
                        "quote": "Queda acreditada la concurrencia continuada y deliberada para falsear la competencia en el mercado nacional.",
                    }
                ],
                "key_points_evidence": [
                    {
                        "point": "Infracción por objeto tipificada en el artículo 1 de la Ley 15/2007.",
                        "quotes": [
                            {
                                "source_field": "content",
                                "quote": "Queda acreditada la concurrencia continuada y deliberada para falsear la competencia en el mercado nacional.",
                            }
                        ],
                    }
                ],
            }
        },
    )
    db_session.add_all([eat_1, call_1])

    # Document 2: Fully analysed vertical agreements decision
    entry_2 = Entry(
        source_id=source.id,
        title="Sentencia TS 128/2026 sobre Acuerdos Verticales de Distribución Exclusiva",
        url="https://cnmc.example.org/ts-128-2026",
        canonical_url="https://cnmc.example.org/ts-128-2026",
        content=(
            "La Sala de lo Contencioso-Administrativo desestima las pretensiones de nulidad contractual. "
            "Las restricciones verticales impuestas garantizaban la integridad de la red de servicio técnico oficial."
        ),
        excerpt="Desestimación de recurso sobre distribución exclusiva...",
        author="Magistrado Ponente D. Alberto Martínez",
        published_at=now - timedelta(days=5),
        captured_at=now - timedelta(days=5),
        content_type="judicial_decision",
        language="es",
    )
    db_session.add(entry_2)
    db_session.flush()

    content_hash_2 = compute_analysis_input_hash(entry_2)
    analysis_2 = EntryAnalysis(
        entry_id=entry_2.id,
        matrix_id=matrix.id,
        pipeline_version="v7",
        status="completed",
        entry_content_hash=content_hash_2,
        matrix_snapshot={"topics": [{"code": "vertical_agreements", "name": "Acuerdos Verticales"}]},
        matrix_snapshot_hash="hash-dossier-2",
        relevance_status="relevant",
        relevance_score=85,
        confidence=0.88,
        summary="El Tribunal Supremo valida restricciones territoriales amparadas en exención por categorías de distribución.",
        key_points=[
            "Compatibilidad de las cláusulas de no competencia con las directrices comunitarias.",
        ],
        completed_at=now - timedelta(days=4),
    )
    db_session.add(analysis_2)
    db_session.flush()

    eat_2 = EntryAnalysisTopic(analysis_id=analysis_2.id, topic_id=topic2.id, is_primary=True)
    call_2 = AnalysisCall(
        entry_analysis_id=analysis_2.id,
        prompt_version_id=prompt_version.id,
        stage="deep_analysis",
        provider="gemini_api",
        model="gemini-3.8-flash",
        status="completed",
        raw_response={
            "result": {
                "summary_evidence": [
                    {
                        "source_field": "content",
                        "quote": "Las restricciones verticales impuestas garantizaban la integridad de la red de servicio técnico oficial.",
                    }
                ],
                "key_points_evidence": [],
            }
        },
    )
    db_session.add_all([eat_2, call_2])

    # Document 3: Unanalysed archive entry
    entry_3 = Entry(
        source_id=source.id,
        title="Circular Informativa 3/2026 sobre Supervisión de Redes de Transporte",
        url="https://cnmc.example.org/circular-3-2026",
        canonical_url="https://cnmc.example.org/circular-3-2026",
        content="Disposiciones reglamentarias sobre peajes de acceso y liquidaciones del sistema de transporte.",
        excerpt="Disposiciones reglamentarias sobre peajes...",
        author="Dirección de Energía",
        published_at=now - timedelta(days=2),
        captured_at=now - timedelta(days=2),
        content_type="article",
        language="es",
    )
    db_session.add(entry_3)
    db_session.commit()

    return {
        "entry_1": entry_1,
        "entry_2": entry_2,
        "entry_3": entry_3,
    }


# ==============================================================================
# 1. Endpoint Core Functionality Tests
# ==============================================================================

def test_01_export_dossier_single_document_success(client: TestClient, setup_dossier_data):
    """1. POST correcto con 1 documento -> HTTP 200 y DOCX válido."""
    doc_id = str(setup_dossier_data["entry_1"].id)
    res = client.post(
        "/api/v1/documents/dossier/export",
        json={"document_ids": [doc_id]},
    )

    assert res.status_code == 200
    assert res.headers["content-type"] == DOCX_MIME_TYPE
    assert "attachment; filename=" in res.headers["content-disposition"]
    assert res.headers["content-disposition"].endswith('.docx"')
    assert res.headers.get("x-content-type-options") == "nosniff"

    docx_doc = docx.Document(io.BytesIO(res.content))
    all_text = "\n".join(p.text for p in docx_doc.paragraphs)
    assert "HITCHINGS" in all_text
    assert "DOCUMENTO 1 DE 1" in all_text
    assert "Resolución CNMC S/0010/24 sobre Cártel de Fabricantes" in all_text


def test_02_export_dossier_multiple_documents_success(client: TestClient, setup_dossier_data):
    """2. POST correcto con varios documentos -> HTTP 200 con capítulos numerados e índice."""
    id1 = str(setup_dossier_data["entry_1"].id)
    id2 = str(setup_dossier_data["entry_2"].id)
    id3 = str(setup_dossier_data["entry_3"].id)

    res = client.post(
        "/api/v1/documents/dossier/export",
        json={"document_ids": [id1, id2, id3]},
    )

    assert res.status_code == 200
    docx_doc = docx.Document(io.BytesIO(res.content))
    all_text = "\n".join(p.text for p in docx_doc.paragraphs)

    assert "DOCUMENTO 1 DE 3" in all_text
    assert "DOCUMENTO 2 DE 3" in all_text
    assert "DOCUMENTO 3 DE 3" in all_text

    # Verifica tabla del índice
    index_table = next(t for t in docx_doc.tables if len(t.columns) == 5)
    assert len(index_table.rows) == 4  # 1 cabecera + 3 documentos


def test_03_export_dossier_preserves_exact_client_order(client: TestClient, setup_dossier_data):
    """3. Se respeta EXACTAMENTE el orden de IDs recibido en la solicitud (orden de entrada es la autoridad)."""
    id1 = str(setup_dossier_data["entry_1"].id)
    id2 = str(setup_dossier_data["entry_2"].id)
    title_1 = "Resolución CNMC S/0010/24 sobre Cártel de Fabricantes"
    title_2 = "Sentencia TS 128/2026 sobre Acuerdos Verticales"

    # Prueba Orden A: [id2, id1]
    res_a = client.post(
        "/api/v1/documents/dossier/export",
        json={"document_ids": [id2, id1]},
    )
    assert res_a.status_code == 200
    doc_a = docx.Document(io.BytesIO(res_a.content))
    text_a = "\n".join(p.text for p in doc_a.paragraphs)

    pos_title_2 = text_a.find(title_2)
    pos_title_1 = text_a.find(title_1)
    assert pos_title_2 != -1 and pos_title_1 != -1
    # Documento 2 debe aparecer antes que Documento 1
    assert pos_title_2 < pos_title_1

    # Índice tabla en Orden A
    index_a = next(t for t in doc_a.tables if len(t.columns) == 5)
    assert "Sentencia TS 128/2026" in index_a.rows[1].cells[1].text
    assert "Resolución CNMC S/0010/24" in index_a.rows[2].cells[1].text

    # Prueba Orden B: [id1, id2]
    res_b = client.post(
        "/api/v1/documents/dossier/export",
        json={"document_ids": [id1, id2]},
    )
    assert res_b.status_code == 200
    doc_b = docx.Document(io.BytesIO(res_b.content))
    text_b = "\n".join(p.text for p in doc_b.paragraphs)

    pos_title_1_b = text_b.find(title_1)
    pos_title_2_b = text_b.find(title_2)
    # Documento 1 debe aparecer antes que Documento 2
    assert pos_title_1_b < pos_title_2_b


def test_04_export_dossier_custom_title(client: TestClient, setup_dossier_data):
    """4. Título personalizado en la solicitud se refleja en portada, metadatos y filename."""
    id1 = str(setup_dossier_data["entry_1"].id)
    custom_title = "Dossier de Cárteles y Competencia 2026"

    res = client.post(
        "/api/v1/documents/dossier/export",
        json={"document_ids": [id1], "title": custom_title},
    )

    assert res.status_code == 200
    assert "dossier-carteles-y-competencia-2026.docx" in res.headers["content-disposition"]

    docx_doc = docx.Document(io.BytesIO(res.content))
    all_text = "\n".join(p.text for p in docx_doc.paragraphs)
    assert custom_title in all_text


def test_05_export_dossier_without_title_uses_default(client: TestClient, setup_dossier_data):
    """5. Sin título (nulo u omitido) utiliza el título por defecto del servicio."""
    id1 = str(setup_dossier_data["entry_1"].id)

    res = client.post(
        "/api/v1/documents/dossier/export",
        json={"document_ids": [id1]},
    )

    assert res.status_code == 200
    assert "dossier-documental-hitchings.docx" in res.headers["content-disposition"]

    docx_doc = docx.Document(io.BytesIO(res.content))
    all_text = "\n".join(p.text for p in docx_doc.paragraphs)
    assert "Dossier documental HITCHINGS" in all_text


def test_06_export_dossier_unanalysed_document(client: TestClient, setup_dossier_data):
    """6. Documento sin análisis se incluye correctamente sin secciones IA ficticias."""
    id3 = str(setup_dossier_data["entry_3"].id)

    res = client.post(
        "/api/v1/documents/dossier/export",
        json={"document_ids": [id3]},
    )

    assert res.status_code == 200
    docx_doc = docx.Document(io.BytesIO(res.content))
    all_text = "\n".join(p.text for p in docx_doc.paragraphs)

    assert "Circular Informativa 3/2026" in all_text
    assert "Documento incorporado sin análisis IA" in all_text
    assert "Expediente en archivo documental original sin análisis automatizado" in all_text
    assert "DICTAMEN JURÍDICO" not in all_text

    index_table = next(t for t in docx_doc.tables if len(t.columns) == 5)
    assert "Sin análisis IA" in index_table.rows[1].cells[4].text


def test_07_export_dossier_mixed_analysed_and_unanalysed(client: TestClient, setup_dossier_data):
    """7. Dossier mixto maneja limpiamente documentos analizados y no analizados."""
    id1 = str(setup_dossier_data["entry_1"].id)
    id3 = str(setup_dossier_data["entry_3"].id)

    res = client.post(
        "/api/v1/documents/dossier/export",
        json={"document_ids": [id1, id3]},
    )

    assert res.status_code == 200
    docx_doc = docx.Document(io.BytesIO(res.content))
    all_text = "\n".join(p.text for p in docx_doc.paragraphs)

    assert "DICTAMEN JURÍDICO Y CALIFICACIÓN DE RELEVANCIA" in all_text
    assert "96/100" in all_text
    assert "Documento incorporado sin análisis IA" in all_text


# ==============================================================================
# 2. Validation & Error Handling Tests
# ==============================================================================

def test_08_export_dossier_nonexistent_document_returns_404(client: TestClient, setup_dossier_data):
    """8. Documento inexistente devuelve HTTP 404 e identifica el ID ausente (sin dossier parcial)."""
    id1 = str(setup_dossier_data["entry_1"].id)
    missing_id = str(uuid.uuid4())

    res = client.post(
        "/api/v1/documents/dossier/export",
        json={"document_ids": [id1, missing_id]},
    )

    assert res.status_code == 404
    detail = res.json()["detail"]
    assert missing_id in detail
    assert "no fueron encontrados" in detail.lower()


def test_09_export_dossier_duplicate_ids_rejected_422(client: TestClient, setup_dossier_data):
    """9. IDs duplicados en la solicitud son rechazados con HTTP 422."""
    id1 = str(setup_dossier_data["entry_1"].id)

    res = client.post(
        "/api/v1/documents/dossier/export",
        json={"document_ids": [id1, id1]},
    )

    assert res.status_code == 422
    assert "duplicados" in str(res.json()).lower()


def test_10_export_dossier_empty_ids_rejected_422(client: TestClient):
    """10. Lista vacía de IDs es rechazada con HTTP 422."""
    res = client.post(
        "/api/v1/documents/dossier/export",
        json={"document_ids": []},
    )

    assert res.status_code == 422


def test_11_export_dossier_more_than_30_ids_rejected_422(client: TestClient):
    """11. Más de 30 identificadores son rechazados con HTTP 422."""
    oversized = [str(uuid.uuid4()) for _ in range(31)]

    res = client.post(
        "/api/v1/documents/dossier/export",
        json={"document_ids": oversized},
    )

    assert res.status_code == 422


def test_12_export_dossier_invalid_uuid_rejected_422(client: TestClient):
    """12. Identificador UUID malformado devuelve HTTP 422."""
    res = client.post(
        "/api/v1/documents/dossier/export",
        json={"document_ids": ["not-a-valid-uuid"]},
    )

    assert res.status_code == 422


def test_13_export_dossier_unauthenticated_returns_401(client: TestClient, setup_dossier_data):
    """13. Usuario no autenticado recibe HTTP 401."""
    id1 = str(setup_dossier_data["entry_1"].id)
    override = app.dependency_overrides.pop(get_current_user, None)
    try:
        res = client.post(
            "/api/v1/documents/dossier/export",
            json={"document_ids": [id1]},
        )
        assert res.status_code == 401
    finally:
        if override:
            app.dependency_overrides[get_current_user] = override


def test_14_export_dossier_content_type_and_disposition_headers(client: TestClient, setup_dossier_data):
    """14. Verificación estricta de Content-Type y Content-Disposition."""
    id1 = str(setup_dossier_data["entry_1"].id)
    res = client.post(
        "/api/v1/documents/dossier/export",
        json={"document_ids": [id1]},
    )

    assert res.status_code == 200
    assert res.headers["content-type"] == DOCX_MIME_TYPE
    assert res.headers["content-disposition"].startswith('attachment; filename="dossier-')
    assert res.headers["content-disposition"].endswith('.docx"')


def test_15_export_dossier_valid_docx_zip_binary(client: TestClient, setup_dossier_data):
    """15. El contenido retornado es un binario ZIP/DOCX válido."""
    id1 = str(setup_dossier_data["entry_1"].id)
    res = client.post(
        "/api/v1/documents/dossier/export",
        json={"document_ids": [id1]},
    )

    assert res.status_code == 200
    assert res.content.startswith(b"PK\x03\x04")
    docx_doc = docx.Document(io.BytesIO(res.content))
    assert len(docx_doc.sections) >= 1


def test_16_export_dossier_contains_actual_document_content(client: TestClient, setup_dossier_data):
    """16. El DOCX contiene el texto real, ficha técnica y evidencias de los documentos."""
    id1 = str(setup_dossier_data["entry_1"].id)
    res = client.post(
        "/api/v1/documents/dossier/export",
        json={"document_ids": [id1]},
    )

    docx_doc = docx.Document(io.BytesIO(res.content))
    all_text = "\n".join(p.text for p in docx_doc.paragraphs)

    assert "La Dirección de Competencia de la CNMC sanciona a cinco mercantiles" in all_text
    assert "Infracción por objeto tipificada en el artículo 1 de la Ley 15/2007" in all_text
    assert "Queda acreditada la concurrencia continuada y deliberada" in all_text


# ==============================================================================
# 3. Read-Only, Isolation & Security Guarantees
# ==============================================================================

def test_17_export_dossier_no_temporary_files_created(client: TestClient, setup_dossier_data):
    """17. No se generan archivos temporales en el directorio de trabajo ni en tempfile."""
    id1 = str(setup_dossier_data["entry_1"].id)
    cwd_before = set(os.listdir("."))
    temp_before = set(os.listdir(tempfile.gettempdir()))

    res = client.post(
        "/api/v1/documents/dossier/export",
        json={"document_ids": [id1]},
    )

    assert res.status_code == 200
    cwd_after = set(os.listdir("."))
    temp_after = set(os.listdir(tempfile.gettempdir()))

    assert cwd_after == cwd_before
    new_temp = temp_after - temp_before
    assert not any(f.endswith(".docx") for f in new_temp)


def test_18_export_dossier_does_not_call_gemini_or_ai(client: TestClient, setup_dossier_data):
    """18. El endpoint no realiza llamadas a modelos de IA ni a Gemini."""
    id1 = str(setup_dossier_data["entry_1"].id)

    with mock.patch("app.providers.ai.base.BaseAIProvider.analyze", side_effect=AssertionError("AI model called!")):
        res = client.post(
            "/api/v1/documents/dossier/export",
            json={"document_ids": [id1]},
        )
        assert res.status_code == 200


def test_19_export_dossier_does_not_mutate_database(client: TestClient, db_session: Session, setup_dossier_data):
    """19. Operación estrictamente de solo lectura sin inserciones ni mutaciones en la BD."""
    id1 = str(setup_dossier_data["entry_1"].id)
    id2 = str(setup_dossier_data["entry_2"].id)

    count_entries_before = db_session.execute(select(func.count(Entry.id))).scalar()
    count_analyses_before = db_session.execute(select(func.count(EntryAnalysis.id))).scalar()
    count_calls_before = db_session.execute(select(func.count(AnalysisCall.id))).scalar()

    res = client.post(
        "/api/v1/documents/dossier/export",
        json={"document_ids": [id1, id2]},
    )
    assert res.status_code == 200

    count_entries_after = db_session.execute(select(func.count(Entry.id))).scalar()
    count_analyses_after = db_session.execute(select(func.count(EntryAnalysis.id))).scalar()
    count_calls_after = db_session.execute(select(func.count(AnalysisCall.id))).scalar()

    assert count_entries_before == count_entries_after
    assert count_analyses_before == count_analyses_after
    assert count_calls_before == count_calls_after


def test_20_existing_individual_export_unaffected(client: TestClient, setup_dossier_data):
    """20. El endpoint individual GET /api/v1/documents/{document_id}/export sigue funcionando."""
    doc_id = str(setup_dossier_data["entry_1"].id)
    res = client.get(f"/api/v1/documents/{doc_id}/export")

    assert res.status_code == 200
    assert res.headers["content-type"] == DOCX_MIME_TYPE
    assert "attachment; filename=" in res.headers["content-disposition"]


def test_21_security_malicious_title_content_disposition(client: TestClient, setup_dossier_data):
    """21. Seguridad: título con caracteres potencialmente problemáticos es sanitizado con seguridad."""
    id1 = str(setup_dossier_data["entry_1"].id)
    malicious_title = '../../etc/passwd"\r\nInjected-Header: evil\r\n'

    res = client.post(
        "/api/v1/documents/dossier/export",
        json={"document_ids": [id1], "title": malicious_title},
    )

    assert res.status_code == 200
    disposition = res.headers["content-disposition"]

    # Previene inyecciones CRLF y path traversal
    assert "\r" not in disposition
    assert "\n" not in disposition
    assert ".." not in disposition
    assert "/" not in disposition
    assert "Injected-Header" not in disposition
    assert disposition.startswith('attachment; filename="dossier-')
    assert disposition.endswith('.docx"')


def test_22_openapi_registry(client: TestClient):
    """22. Verificación de que POST /api/v1/documents/dossier/export aparece en /openapi.json."""
    res = client.get("/openapi.json")
    assert res.status_code == 200
    openapi = res.json()

    path_spec = openapi.get("paths", {}).get("/api/v1/documents/dossier/export", {})
    assert "post" in path_spec
    post_op = path_spec["post"]
    assert "responses" in post_op
    assert "200" in post_op["responses"]
    assert DOCX_MIME_TYPE in post_op["responses"]["200"]["content"]
